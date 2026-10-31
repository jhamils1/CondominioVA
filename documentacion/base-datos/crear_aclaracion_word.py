from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

d=Document(); sec=d.sections[0]; sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.2)
styles=d.styles
styles['Normal'].font.name='Calibri'; styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(6)
for n in ['Title','Heading 1','Heading 2']:
    styles[n].font.name='Calibri'; styles[n].font.color.rgb=RGBColor(0,0,0)
d.add_heading('Comparación del diagrama con el proyecto',0)
d.add_paragraph('El diagrama de referencia y la base de datos definida en los modelos actuales no son iguales. Comparten algunas entidades, pero difieren en la representación de personas, relaciones y módulos. El XMI generado representa el código actual; la comparación no modifica la base de datos.')
d.add_heading('Diferencias entre entidades y relaciones',level=1)
rows=[
('Propietario','Se utiliza Persona con tipo P; no hay una tabla independiente de propietarios.'),
('Residente y tipo_residente','No existen esas tablas. Inquilino, Familiares y Visitante heredan de Persona. Inquilino se relaciona con un propietario y Familiares con una persona relacionada.'),
('propietario/mascotas','No existe la tabla intermedia. Cada Mascota apunta directamente a una Persona, que puede tener varias mascotas.'),
('vehiculo/propietario','No existe la tabla intermedia. Cada Vehiculo apunta directamente a una Persona, que puede tener varios vehículos.'),
('Usuario, rol y permiso','Se usan los modelos User, Group y Permission de Django y sus tablas intermedias.'),
('Rol de usuario','La base permite varios grupos por usuario. La API asigna un solo grupo y muestra el primero como rol.'),
('Cargo y empleado','Se conserva la relación: un cargo puede tener varios empleados.'),
('Propietario, contrato y unidad','Se conserva la estructura, pero el propietario de contrato referencia Persona.'),
('Tipo y unidad','No existe la tabla tipo. Unidad.tipo_unidad identifica apartamento, casa, local comercial o estacionamiento.'),
('Bloque y unidad','Se conserva la relación: un bloque puede contener varias unidades.'),
('Factura, detalle_factura y servicio','No existen como modelos persistentes. Finanzas contiene contrato, expensa y multa, que no son equivalentes directos.'),
('Nomina y detalle de nomina','No están implementadas como tablas. Existen Empleado.sueldo y Cargo.salario_base.')]
t=d.add_table(rows=1, cols=2); t.autofit=False; t.columns[0].width=Cm(5.5); t.columns[1].width=Cm(10)
t.style='Table Grid'; head=t.rows[0]; head.cells[0].text='En el diagrama'; head.cells[1].text='En el proyecto actual'
trPr=head._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:tblHeader'))
for c in head.cells:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DCEAF3'); c._tc.get_or_add_tcPr().append(sh)
    for r in c.paragraphs[0].runs: r.bold=True
for a,b in rows:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b
    for cell in cells:
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(4)
            for r in p.runs: r.font.size=Pt(10)
d.add_heading('Diferencias de cardinalidad',level=1)
for text in [
'En la imagen un propietario tiene uno o más contratos. En el proyecto, una persona puede tener cero o muchos contratos.',
'En la imagen propietario y empleado necesitan un usuario asociado. En el proyecto, cada persona o empleado puede tener cero o un usuario.',
'En la imagen las tablas intermedias permiten asociaciones entre varios propietarios y varias mascotas o vehículos. En el proyecto, cada mascota y cada vehículo tiene una sola persona asociada.'
]: d.add_paragraph(text,style='List Bullet')
d.add_heading('Entidades adicionales del proyecto',level=1)
d.add_paragraph('Visita, AreasComunes, ReservaAreaComun, ObjetoPerdido, incidente, expensa, multa y LecturaPlaca no aparecen en la imagen. También existen tablas auxiliares de Django, JWT y notificaciones.')
d.add_heading('Campos y representación UML',level=1)
d.add_paragraph('El contrato actual utiliza fecha_contrato, cuota_mensual, costo_compra y contrato_PDF. No tiene los campos impuesto, monto, fecha_inicio y fecha_fin que muestra la imagen.')
d.add_paragraph('El proyecto declara herencia desde Inquilino, Familiares y Visitante hacia Persona; se representa con triángulos UML. Las demás claves foráneas se representan como asociaciones. La eliminación en cascada no implica por sí sola una composición UML.')
d.add_heading('Fuentes revisadas',level=1)
d.add_paragraph('Condominio-Backend: administracion/models.py, residencial/models.py, residencial/modelsVehiculo.py, finanzas/models.py, seguridad_IA/models.py y serializadores de usuarios y roles.')
out=Path(__file__).parent/'comparacion_diagrama_base_datos.docx'; d.save(out); print(out)
